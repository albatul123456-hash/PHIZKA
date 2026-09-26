import os
import io
import math
import shutil
import subprocess
import tempfile
from datetime import datetime

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from scipy.signal import find_peaks
from scipy.stats import linregress

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# إعداد PHIZKA
# ============================================================

st.set_page_config(
    page_title="PHIZKA | فيزكا العلمية",
    page_icon="⚛️",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# تنسيق الواجهة
# ============================================================

st.markdown(
    """
    <style>
    .main-title {
        text-align:center;
        font-size:46px;
        font-weight:800;
        margin-bottom:0px;
    }

    .sub-title {
        text-align:center;
        font-size:20px;
        opacity:0.75;
        margin-bottom:30px;
    }

    .phizka-card {
        padding:22px;
        border:1px solid rgba(128,128,128,0.25);
        border-radius:18px;
        margin-bottom:15px;
    }

    div[data-testid="stMetric"] {
        border:1px solid rgba(128,128,128,0.25);
        padding:12px;
        border-radius:14px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# أدوات البيانات
# ============================================================

def safe_numeric_dataframe(df, columns):

    clean = df.copy()

    for col in columns:
        clean[col] = pd.to_numeric(
            clean[col],
            errors="coerce"
        )

    return clean.dropna(
        subset=columns
    )


def find_gnuplot():

    for cmd in ("gnuplot", "wgnuplot"):

        path = shutil.which(cmd)

        if path:
            return path

    possible_paths = [
        r"C:\Program Files\gnuplot\bin\gnuplot.exe",
        r"C:\Program Files\gnuplot\bin\wgnuplot.exe",
        r"C:\gnuplot\bin\gnuplot.exe",
        r"C:\gnuplot\bin\wgnuplot.exe",
    ]

    for path in possible_paths:

        if os.path.exists(path):
            return path

    return None


gnuplot_path = find_gnuplot()


# ============================================================
# Excel
# ============================================================

def create_excel_file(dataframes, results=None):

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        for sheet_name, dataframe in dataframes.items():

            dataframe.to_excel(
                writer,
                sheet_name=sheet_name[:31],
                index=False
            )

        if results:

            results_df = pd.DataFrame(
                {
                    "Result": list(results.keys()),
                    "Value": list(results.values())
                }
            )

            results_df.to_excel(
                writer,
                sheet_name="Results",
                index=False
            )

    output.seek(0)

    return output.getvalue()


# ============================================================
# الرسم إلى صورة لاستخدامه داخل Word
# ============================================================

def figure_to_bytes(fig):

    image_stream = io.BytesIO()

    fig.savefig(
        image_stream,
        format="png",
        dpi=180,
        bbox_inches="tight"
    )

    image_stream.seek(0)

    return image_stream


# ============================================================
# أدوات Word
# ============================================================

def set_paragraph_rtl(paragraph):

    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_arabic_heading(document, text, level=1):

    heading = document.add_heading(
        text,
        level=level
    )

    set_paragraph_rtl(heading)

    return heading


def add_arabic_paragraph(document, text, bold=False):

    paragraph = document.add_paragraph()

    set_paragraph_rtl(paragraph)

    run = paragraph.add_run(
        str(text)
    )

    run.bold = bold
    run.font.size = Pt(11)

    return paragraph


def add_dataframe_to_word(document, dataframe):

    if dataframe.empty:

        add_arabic_paragraph(
            document,
            "لا توجد بيانات."
        )

        return

    table = document.add_table(
        rows=1,
        cols=len(dataframe.columns)
    )

    table.style = "Table Grid"

    for index, column in enumerate(
        dataframe.columns
    ):

        table.rows[0].cells[index].text = str(
            column
        )

    for row in dataframe.itertuples(
        index=False,
        name=None
    ):

        cells = table.add_row().cells

        for index, value in enumerate(row):

            if isinstance(
                value,
                (float, np.floating)
            ):

                if np.isfinite(value):

                    cells[index].text = (
                        f"{value:.6g}"
                    )

                else:

                    cells[index].text = ""

            else:

                cells[index].text = str(
                    value
                )


def create_experiment_report(
    experiment_title,
    objective,
    theory,
    dataframe,
    formulas,
    calculations,
    results,
    figure=None,
    analysis_text=None
):

    document = Document()

    title = document.add_heading(
        "PHIZKA Scientific Platform",
        0
    )

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(
        "فيزكا العلمية"
    )

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_arabic_heading(
        document,
        experiment_title,
        level=1
    )

    add_arabic_paragraph(
        document,
        f"تاريخ إنشاء التقرير: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )

    add_arabic_heading(
        document,
        "الهدف من التجربة",
        level=2
    )

    add_arabic_paragraph(
        document,
        objective
    )

    if theory:

        add_arabic_heading(
            document,
            "الأساس النظري",
            level=2
        )

        add_arabic_paragraph(
            document,
            theory
        )

    add_arabic_heading(
        document,
        "القراءات التجريبية",
        level=2
    )

    add_dataframe_to_word(
        document,
        dataframe
    )

    add_arabic_heading(
        document,
        "القوانين المستخدمة",
        level=2
    )

    for formula in formulas:

        add_arabic_paragraph(
            document,
            formula,
            bold=True
        )

    add_arabic_heading(
        document,
        "خطوات الحساب والتعويض",
        level=2
    )

    for calculation in calculations:

        add_arabic_paragraph(
            document,
            calculation
        )

    add_arabic_heading(
        document,
        "النتائج النهائية",
        level=2
    )

    for name, value in results.items():

        add_arabic_paragraph(
            document,
            f"{name}: {value}"
        )

    if analysis_text:

        add_arabic_heading(
            document,
            "التحليل",
            level=2
        )

        add_arabic_paragraph(
            document,
            analysis_text
        )

    if figure is not None:

        add_arabic_heading(
            document,
            "الرسم البياني",
            level=2
        )

        image_stream = figure_to_bytes(
            figure
        )

        document.add_picture(
            image_stream,
            width=Inches(6.2)
        )

    add_arabic_heading(
        document,
        "الخلاصة",
        level=2
    )

    add_arabic_paragraph(
        document,
        "تم إجراء الحسابات آليًا باستخدام القراءات "
        "والقيم المدخلة في منصة PHIZKA. عند تعديل "
        "القراءات أو قيم التعويض تتغير النتائج والتقرير "
        "وفقًا للقيم الجديدة."
    )

    output = io.BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output.getvalue()


# ============================================================
# Linear Regression
# ============================================================

def calculate_regression(x, y):

    x = np.asarray(
        x,
        dtype=float
    )

    y = np.asarray(
        y,
        dtype=float
    )

    valid = (
        np.isfinite(x)
        &
        np.isfinite(y)
    )

    x = x[valid]
    y = y[valid]

    if len(x) < 2:

        return None

    if np.allclose(
        x,
        x[0]
    ):

        return None

    result = linregress(
        x,
        y
    )

    return {
        "slope": float(result.slope),
        "intercept": float(result.intercept),
        "r2": float(result.rvalue ** 2),
    }


# ============================================================
# رأس المنصة
# ============================================================

st.markdown(
    '<div class="main-title">⚛️ PHIZKA</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="sub-title">'
    'فيزكا العلمية | PHIZKA Scientific Platform'
    '</div>',
    unsafe_allow_html=True,
)


# ============================================================
# القائمة
# ============================================================

page = st.sidebar.radio(
    "⚛️ أقسام فيزكا",
    [
        "🏠 الرئيسية",
        "📚 المقررات العملية",
        "🧮 الحاسبة الفيزيائية",
        "📊 تحليل البيانات",
        "🧪 ورشة تحليل البيانات الفيزيائية",
        "🧊 Gnuplot 3D",
        "🎓 مشروع التخرج والبحث العلمي",
    ],
)

st.sidebar.divider()

if gnuplot_path:

    st.sidebar.success(
        "Gnuplot جاهز ✅"
    )

else:

    st.sidebar.info(
        "Gnuplot غير مكتشف"
    )


# ============================================================
# الرئيسية
# ============================================================

if page == "🏠 الرئيسية":

    st.header(
        "مرحبًا بك في فيزكا العلمية 🔬"
    )

    st.write(
        """
        منصة تعليمية علمية تفاعلية لمساعدة طلبة الفيزياء
        في التجارب العملية والحسابات والرسوم البيانية
        وتحليل البيانات وإصدار التقارير العلمية.
        """
    )

    c1, c2, c3 = st.columns(3)

    with c1:

        st.markdown(
            """
            <div class="phizka-card">
            <h3>📚 المقررات العملية</h3>
            تجارب الفيزياء مع القراءات والقوانين
            والحساب والرسم والتحليل.
            </div>
            """,
            unsafe_allow_html=True,
        )

    with c2:

        st.markdown(
            """
            <div class="phizka-card">
            <h3>🧮 الحاسبة الفيزيائية</h3>
            أدوات حسابية سريعة مع التعويض
            داخل القوانين.
            </div>
            """,
            unsafe_allow_html=True,
        )

    with c3:

        st.markdown(
            """
            <div class="phizka-card">
            <h3>📊 التحليل والتقارير</h3>
            Excel وCSV وWord وR²
            والرسوم العلمية.
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.divider()

    st.subheader(
        "📚 المقررات المتوفرة"
    )

    a, b = st.columns(2)

    with a:

        st.info(
            "🧊 فيزياء الجوامد 1\n\n"
            "• Bragg's Law for NaCl\n\n"
            "• الخلية الشمسية"
        )

    with b:

        st.info(
            "⚛️ فيزياء نووية 1\n\n"
            "• تحديد منطقة التشغيل لعداد جايجر"
        )


# ============================================================
# المقررات العملية
# ============================================================

elif page == "📚 المقررات العملية":

    st.header(
        "📚 المقررات العملية"
    )

    course = st.selectbox(
        "اختاري المقرر",
        [
            "🧊 فيزياء الجوامد 1",
            "⚛️ فيزياء نووية 1",
        ],
    )


    # ========================================================
    # جوامد 1
    # ========================================================

    if course == "🧊 فيزياء الجوامد 1":

        st.header(
            "🧊 فيزياء الجوامد 1"
        )

        experiment = st.selectbox(
            "اختاري التجربة",
            [
                "1️⃣ Bragg's Law for NaCl",
                "2️⃣ الخلية الشمسية | Solar Cell",
            ],
        )


        # ====================================================
        # BRAGG
        # ====================================================

        if experiment == "1️⃣ Bragg's Law for NaCl":

            st.subheader(
                "🔬 إيجاد المسافة البينية لبلورة NaCl"
            )

            st.write(
                "الهدف: إيجاد المسافة البينية بين "
                "مستويات بلورة NaCl باستخدام قانون براج."
            )

            st.latex(
                r"n\lambda=2d\sin(\theta)"
            )

            st.latex(
                r"d=\frac{n\lambda}{2\sin(\theta)}"
            )

            wavelength = st.number_input(
                "λ",
                min_value=0.0001,
                value=1.456,
                step=0.001,
                format="%.3f",
                key="bragg_wavelength"
            )

            st.header(
                "📋 القراءات"
            )

            bragg_default = pd.DataFrame(
                {
                    "Theta θ (degree)": [
                        3, 3.5, 4, 4.5, 5,
                        5.5, 6, 6.5, 7, 7.5,
                        8, 8.5, 9, 9.5, 10,
                        10.5, 11, 11.5, 12, 12.5,
                        13, 13.5, 14, 14.5, 15,
                        15.5, 16, 16.5, 17, 17.5,
                        18, 18.5, 19, 19.5, 20,
                        20.5, 21, 21.5, 22, 22.5,
                        23
                    ],

                    "Intensity I": [
                        8, 15, 514, 677, 655,
                        628, 503, 947, 593, 463,
                        286, 226, 209, 156, 116,
                        130, 113, 97, 89, 80,
                        252, 76, 72, 569, 70,
                        77, 46, 35, 50, 49,
                        36, 24, 32, 62, 27,
                        27, 24, 36, 76, 30,
                        21
                    ]
                }
            )

            bragg_edit = st.data_editor(
                bragg_default,
                num_rows="dynamic",
                use_container_width=True,
                key="bragg_table"
            )

            bragg = safe_numeric_dataframe(
                bragg_edit,
                [
                    "Theta θ (degree)",
                    "Intensity I"
                ]
            )

            bragg = bragg.sort_values(
                "Theta θ (degree)"
            )

            bragg_fig = None
            peaks = np.array([], dtype=int)

            if len(bragg) >= 2:

                theta = bragg[
                    "Theta θ (degree)"
                ].to_numpy(float)

                intensity = bragg[
                    "Intensity I"
                ].to_numpy(float)

                prominence = st.number_input(
                    "Prominence",
                    min_value=1.0,
                    value=50.0,
                    step=5.0,
                    key="bragg_prominence"
                )

                peaks, _ = find_peaks(
                    intensity,
                    prominence=prominence
                )

                bragg_fig, ax = plt.subplots(
                    figsize=(11, 5)
                )

                ax.plot(
                    theta,
                    intensity,
                    marker="o",
                    label="Experimental Data"
                )

                if len(peaks) > 0:

                    ax.scatter(
                        theta[peaks],
                        intensity[peaks],
                        c="red",
                        s=90,
                        label="Detected Peaks"
                    )

                    for peak_number, idx in enumerate(
                        peaks,
                        start=1
                    ):

                        ax.annotate(
                            f"P{peak_number}",
                            (
                                theta[idx],
                                intensity[idx]
                            ),
                            xytext=(5, 8),
                            textcoords="offset points"
                        )

                ax.set_xlabel(
                    "Theta θ (degree)"
                )

                ax.set_ylabel(
                    "Intensity I"
                )

                ax.set_title(
                    "Bragg Diffraction - NaCl"
                )

                ax.grid(
                    True,
                    alpha=0.3
                )

                ax.legend()

                bragg_fig.tight_layout()

                st.pyplot(
                    bragg_fig
                )

            st.header(
                "🧮 الحساب والتعويض"
            )

            c1, c2, c3 = st.columns(3)

            with c1:

                calc_n = st.number_input(
                    "n",
                    min_value=1,
                    value=1,
                    step=1,
                    key="calc_n"
                )

            with c2:

                calc_lambda = st.number_input(
                    "λ المستخدمة في الحساب",
                    min_value=0.0001,
                    value=float(wavelength),
                    step=0.001,
                    format="%.3f",
                    key="calc_lambda"
                )

            with c3:

                default_angle = 6.5

                if (
                    len(bragg) >= 2
                    and
                    len(peaks) > 0
                ):

                    default_angle = float(
                        theta[peaks[0]]
                    )

                calc_theta = st.number_input(
                    "θ المستخدمة في الحساب",
                    value=default_angle,
                    step=0.1,
                    key="calc_theta"
                )

            sin_theta = np.sin(
                np.deg2rad(calc_theta)
            )

            d = None

            if np.isclose(
                sin_theta,
                0
            ):

                st.error(
                    "لا يمكن الحساب لأن sin(θ) = 0."
                )

            else:

                d = (
                    calc_n
                    *
                    calc_lambda
                ) / (
                    2
                    *
                    sin_theta
                )

                st.latex(
                    r"d="
                    r"\frac{n\lambda}"
                    r"{2\sin\theta}"
                )

                st.latex(
                    rf"d="
                    rf"\frac{{"
                    rf"({calc_n})"
                    rf"({calc_lambda:.3f})"
                    rf"}}"
                    rf"{{"
                    rf"2\sin({calc_theta:.2f}^\circ)"
                    rf"}}"
                )

                st.latex(
                    rf"d={d:.4f}"
                )

            # --------------------------------------------
            # حساب d للقمم
            # --------------------------------------------

            peak_results = []

            if (
                len(bragg) >= 2
                and
                len(peaks) > 0
            ):

                st.subheader(
                    "📌 حساب d للقمم المكتشفة"
                )

                number_peaks = min(
                    len(peaks),
                    4
                )

                selected_peaks = peaks[:number_peaks]

                for i, idx in enumerate(
                    selected_peaks,
                    start=1
                ):

                    peak_angle = float(
                        theta[idx]
                    )

                    peak_order = st.number_input(
                        f"رتبة القمة P{i}",
                        min_value=1,
                        value=i,
                        step=1,
                        key=f"bragg_order_{i}"
                    )

                    peak_sin = np.sin(
                        np.deg2rad(
                            peak_angle
                        )
                    )

                    if not np.isclose(
                        peak_sin,
                        0
                    ):

                        peak_d = (
                            peak_order
                            *
                            wavelength
                        ) / (
                            2
                            *
                            peak_sin
                        )

                        peak_results.append(
                            {
                                "Peak": f"P{i}",
                                "n": peak_order,
                                "Theta": peak_angle,
                                "Intensity": float(
                                    intensity[idx]
                                ),
                                "d": peak_d
                            }
                        )

                if peak_results:

                    peak_df = pd.DataFrame(
                        peak_results
                    )

                    st.dataframe(
                        peak_df,
                        use_container_width=True,
                        hide_index=True
                    )

                    d_average = float(
                        peak_df["d"].mean()
                    )

                    st.latex(
                        r"d_{\mathrm{avg}}="
                        r"\frac{\sum d_i}{N}"
                    )

                    st.success(
                        f"متوسط d = {d_average:.4f}"
                    )

            # --------------------------------------------
            # تصدير Bragg
            # --------------------------------------------

            st.divider()

            st.subheader(
                "📤 تصدير تجربة Bragg"
            )

            bragg_results = {
                "n المستخدم في الحساب": calc_n,
                "λ": f"{calc_lambda:.6g}",
                "θ": f"{calc_theta:.6g} degree",
                "d": (
                    f"{d:.6g}"
                    if d is not None
                    else "غير معرف"
                )
            }

            if peak_results:

                bragg_results[
                    "متوسط d للقمم"
                ] = f"{d_average:.6g}"

            bragg_calculations = [
                "قانون براج: nλ = 2d sin(θ)",
                "بإعادة ترتيب القانون: "
                "d = nλ / [2 sin(θ)]",
                (
                    f"بالتعويض: d = "
                    f"({calc_n} × {calc_lambda:.6g}) "
                    f"/ [2 × sin({calc_theta:.6g}°)]"
                ),
            ]

            if d is not None:

                bragg_calculations.append(
                    f"الناتج: d = {d:.6g}"
                )

            if peak_results:

                for result in peak_results:

                    bragg_calculations.append(
                        f"{result['Peak']}: "
                        f"d = "
                        f"({result['n']} × {wavelength:.6g}) "
                        f"/ [2 × sin({result['Theta']:.6g}°)] "
                        f"= {result['d']:.6g}"
                    )

                bragg_calculations.append(
                    f"متوسط المسافة البينية = "
                    f"{d_average:.6g}"
                )

            bragg_csv = bragg.to_csv(
                index=False
            ).encode(
                "utf-8-sig"
            )

            bragg_excel_frames = {
                "Readings": bragg
            }

            if peak_results:

                bragg_excel_frames[
                    "Peak Results"
                ] = pd.DataFrame(
                    peak_results
                )

            bragg_excel = create_excel_file(
                bragg_excel_frames,
                bragg_results
            )

            bragg_word = create_experiment_report(
                experiment_title=(
                    "تجربة قانون براج لبلورة NaCl"
                ),
                objective=(
                    "إيجاد المسافة البينية بين مستويات "
                    "بلورة كلوريد الصوديوم باستخدام قانون براج."
                ),
                theory=(
                    "تظهر قمم الحيود عندما يتحقق شرط براج "
                    "بين الطول الموجي وزاوية الحيود "
                    "والمسافة البينية للمستويات البلورية."
                ),
                dataframe=bragg,
                formulas=[
                    "nλ = 2d sin(θ)",
                    "d = nλ / [2 sin(θ)]",
                    "d_avg = Σdᵢ / N"
                ],
                calculations=bragg_calculations,
                results=bragg_results,
                figure=bragg_fig
            )

            e1, e2, e3 = st.columns(3)

            with e1:

                st.download_button(
                    "📄 CSV",
                    bragg_csv,
                    "PHIZKA_Bragg.csv",
                    "text/csv"
                )

            with e2:

                st.download_button(
                    "📊 Excel",
                    bragg_excel,
                    "PHIZKA_Bragg.xlsx",
                    (
                        "application/vnd.openxmlformats-"
                        "officedocument.spreadsheetml.sheet"
                    )
                )

            with e3:

                st.download_button(
                    "📝 تقرير Word كامل",
                    bragg_word,
                    "PHIZKA_Bragg_Report.docx",
                    (
                        "application/vnd.openxmlformats-"
                        "officedocument.wordprocessingml.document"
                    )
                )

            if bragg_fig is not None:

                plt.close(
                    bragg_fig
                )


        # ====================================================
        # SOLAR CELL
        # ====================================================

        elif experiment == "2️⃣ الخلية الشمسية | Solar Cell":

            st.header(
                "☀️ الخلية الشمسية"
            )

            st.write(
                "الهدف: دراسة تغير القدرة الكهربائية "
                "مع مقاومة الحمل وتحديد مقاومة الحمل "
                "المثلى ونقطة القدرة العظمى."
            )

            st.latex(
                r"I=\frac{V}{R}"
            )

            st.latex(
                r"P=VI=\frac{V^2}{R}"
            )

            st.latex(
                r"I_0=\sqrt{\frac{P_0}{R_0}}"
            )

            st.latex(
                r"\eta="
                r"\frac{P_0}{P_{in}}\times100"
            )

            solar_default = pd.DataFrame(
                {
                    "R (Ω)": [
                        100, 200, 300, 400, 500,
                        600, 700, 800, 900, 1000
                    ],

                    "V (V)": [
                        1.78, 3.20, 4.56, 5.75,
                        6.97, 7.82, 8.45, 8.82,
                        9.16, 9.38
                    ],

                    "I measured (mA)": [
                        15.10, 14.61, 14.33, 13.79,
                        13.41, 12.64, 11.78, 10.33,
                        9.92, 9.19
                    ]
                }
            )

            st.subheader(
                "📋 القراءات"
            )

            solar_edit = st.data_editor(
                solar_default,
                num_rows="dynamic",
                use_container_width=True,
                key="solar_table"
            )

            solar = safe_numeric_dataframe(
                solar_edit,
                [
                    "R (Ω)",
                    "V (V)",
                    "I measured (mA)"
                ]
            )

            solar = solar[
                solar["R (Ω)"] > 0
            ].copy()

            calculation_method = st.radio(
                "طريقة حساب القدرة",
                [
                    "من التيار المقاس: P = VI",
                    "من المقاومة والجهد: P = V²/R"
                ],
                horizontal=True
            )

            if calculation_method == "من التيار المقاس: P = VI":

                solar["P (mW)"] = (
                    solar["V (V)"]
                    *
                    solar["I measured (mA)"]
                )

            else:

                solar["I calculated (mA)"] = (
                    solar["V (V)"]
                    /
                    solar["R (Ω)"]
                    *
                    1000
                )

                solar["P (mW)"] = (
                    solar["V (V)"] ** 2
                    /
                    solar["R (Ω)"]
                    *
                    1000
                )

            st.dataframe(
                solar,
                use_container_width=True,
                hide_index=True
            )

            solar_fig = None

            if len(solar) > 0:

                R_values = solar[
                    "R (Ω)"
                ].to_numpy(float)

                P_values = solar[
                    "P (mW)"
                ].to_numpy(float)

                max_index = int(
                    np.argmax(P_values)
                )

                detected_r0 = float(
                    R_values[max_index]
                )

                detected_p0 = float(
                    P_values[max_index]
                )

                solar_fig, ax = plt.subplots(
                    figsize=(11, 5)
                )

                ax.plot(
                    R_values,
                    P_values,
                    marker="o",
                    label="Experimental Data"
                )

                ax.scatter(
                    [detected_r0],
                    [detected_p0],
                    c="red",
                    s=120,
                    label="Maximum Power"
                )

                ax.set_xlabel(
                    "Resistance R (Ω)"
                )

                ax.set_ylabel(
                    "Power P (mW)"
                )

                ax.set_title(
                    "Solar Cell - Power vs Resistance"
                )

                ax.grid(
                    True,
                    alpha=0.3
                )

                ax.legend()

                solar_fig.tight_layout()

                st.pyplot(
                    solar_fig
                )

                st.header(
                    "🧮 الحساب والتعويض"
                )

                c1, c2 = st.columns(2)

                with c1:

                    r0 = st.number_input(
                        "R₀ (Ω)",
                        min_value=0.000001,
                        value=detected_r0,
                        step=10.0,
                        key="solar_r0"
                    )

                with c2:

                    p0_mw = st.number_input(
                        "P₀ (mW)",
                        min_value=0.0,
                        value=detected_p0,
                        step=0.1,
                        key="solar_p0"
                    )

                p0_w = (
                    p0_mw / 1000
                )

                i0 = math.sqrt(
                    p0_w / r0
                )

                st.latex(
                    r"I_0="
                    r"\sqrt{\frac{P_0}{R_0}}"
                )

                st.latex(
                    rf"I_0="
                    rf"\sqrt{{"
                    rf"\frac{{{p0_w:.6f}}}"
                    rf"{{{r0:.3f}}}"
                    rf"}}"
                )

                st.latex(
                    rf"I_0="
                    rf"{i0:.6f}"
                    rf"\,\mathrm{{A}}"
                )

                st.latex(
                    rf"I_0="
                    rf"{i0*1000:.3f}"
                    rf"\,\mathrm{{mA}}"
                )

                pin = st.number_input(
                    "Pᵢₙ (W) لحساب الكفاءة",
                    min_value=0.0,
                    value=0.0,
                    step=0.01,
                    key="solar_pin"
                )

                efficiency = None

                if pin > 0:

                    efficiency = (
                        p0_w / pin
                    ) * 100

                    st.latex(
                        r"\eta="
                        r"\frac{P_0}{P_{in}}\times100"
                    )

                    st.latex(
                        rf"\eta="
                        rf"\frac{{{p0_w:.6f}}}"
                        rf"{{{pin:.6f}}}"
                        rf"\times100"
                    )

                    st.success(
                        f"η = {efficiency:.3f}%"
                    )

                solar_results = {
                    "R₀": f"{r0:.6g} Ω",
                    "P₀": f"{p0_mw:.6g} mW",
                    "P₀ بوحدة W": f"{p0_w:.6g} W",
                    "I₀": f"{i0:.6g} A",
                    "I₀ بوحدة mA": f"{i0*1000:.6g} mA",
                }

                if pin > 0:

                    solar_results[
                        "Pᵢₙ"
                    ] = f"{pin:.6g} W"

                    solar_results[
                        "الكفاءة η"
                    ] = f"{efficiency:.6g}%"

                solar_calculations = []

                for row_number, row in enumerate(
                    solar.itertuples(
                        index=False
                    ),
                    start=1
                ):

                    resistance = float(
                        solar.iloc[
                            row_number - 1
                        ]["R (Ω)"]
                    )

                    voltage = float(
                        solar.iloc[
                            row_number - 1
                        ]["V (V)"]
                    )

                    power = float(
                        solar.iloc[
                            row_number - 1
                        ]["P (mW)"]
                    )

                    if calculation_method == "من التيار المقاس: P = VI":

                        current = float(
                            solar.iloc[
                                row_number - 1
                            ]["I measured (mA)"]
                        )

                        solar_calculations.append(
                            f"القراءة {row_number}: "
                            f"P = VI = "
                            f"({voltage:.6g}) "
                            f"× ({current:.6g} mA) "
                            f"= {power:.6g} mW"
                        )

                    else:

                        solar_calculations.append(
                            f"القراءة {row_number}: "
                            f"P = V²/R = "
                            f"({voltage:.6g})² / "
                            f"({resistance:.6g}) "
                            f"= {power:.6g} mW"
                        )

                solar_calculations.extend(
                    [
                        (
                            f"أكبر قدرة مقاسة هي "
                            f"P₀ = {p0_mw:.6g} mW "
                            f"عند R₀ = {r0:.6g} Ω."
                        ),

                        (
                            f"تحويل القدرة إلى W: "
                            f"P₀ = {p0_mw:.6g}/1000 "
                            f"= {p0_w:.6g} W."
                        ),

                        (
                            f"I₀ = √(P₀/R₀) = "
                            f"√({p0_w:.6g}/{r0:.6g}) "
                            f"= {i0:.6g} A "
                            f"= {i0*1000:.6g} mA."
                        ),
                    ]
                )

                if pin > 0:

                    solar_calculations.append(
                        f"η = (P₀/Pᵢₙ) × 100 = "
                        f"({p0_w:.6g}/{pin:.6g}) × 100 "
                        f"= {efficiency:.6g}%."
                    )

                solar_csv = solar.to_csv(
                    index=False
                ).encode(
                    "utf-8-sig"
                )

                solar_excel = create_excel_file(
                    {
                        "Readings": solar
                    },
                    solar_results
                )

                solar_word = create_experiment_report(
                    experiment_title=(
                        "تجربة الخلية الشمسية"
                    ),
                    objective=(
                        "دراسة تغير القدرة الكهربائية "
                        "مع مقاومة الحمل وتحديد مقاومة الحمل "
                        "المثلى ونقطة القدرة العظمى."
                    ),
                    theory=(
                        "تحول الخلية الشمسية الطاقة الضوئية "
                        "إلى طاقة كهربائية. تتغير القدرة "
                        "المستخرجة مع مقاومة الحمل وتوجد نقطة "
                        "تكون عندها القدرة المقاسة عظمى."
                    ),
                    dataframe=solar,
                    formulas=[
                        "I = V / R",
                        "P = VI",
                        "P = V² / R",
                        "I₀ = √(P₀ / R₀)",
                        "η = (P₀ / Pᵢₙ) × 100"
                    ],
                    calculations=solar_calculations,
                    results=solar_results,
                    figure=solar_fig
                )

                e1, e2, e3 = st.columns(3)

                with e1:

                    st.download_button(
                        "📄 CSV",
                        solar_csv,
                        "PHIZKA_Solar_Cell.csv",
                        "text/csv"
                    )

                with e2:

                    st.download_button(
                        "📊 Excel",
                        solar_excel,
                        "PHIZKA_Solar_Cell.xlsx",
                        (
                            "application/vnd.openxmlformats-"
                            "officedocument.spreadsheetml.sheet"
                        )
                    )

                with e3:

                    st.download_button(
                        "📝 تقرير Word كامل",
                        solar_word,
                        "PHIZKA_Solar_Cell_Report.docx",
                        (
                            "application/vnd.openxmlformats-"
                            "officedocument.wordprocessingml.document"
                        )
                    )

                if solar_fig is not None:

                    plt.close(
                        solar_fig
                    )


    # ========================================================
    # نووية 1
    # ========================================================

    elif course == "⚛️ فيزياء نووية 1":

        st.header(
            "⚛️ فيزياء نووية 1"
        )

        experiment = st.selectbox(
            "اختاري التجربة",
            [
                "☢️ تحديد منطقة التشغيل لعداد جايجر"
            ]
        )

        st.header(
            "☢️ تحديد منطقة التشغيل لعداد جايجر"
        )

        st.subheader(
            "🎯 الهدف"
        )

        st.write(
            """
            • رسم منحنى التشغيل  
            • تعيين منطقة الهضبة  
            • تحديد جهد الانطلاق Vs  
            • تحديد V₁ وV₂  
            • حساب طول منطقة الهضبة  
            • حساب ميل منطقة الهضبة  
            • تحديد جهد التشغيل المناسب
            """
        )

        st.warning(
            "عند حدوث ارتفاع كبير ومفاجئ في معدل العد "
            "يجب التوقف عن زيادة الجهد."
        )

        st.latex(
            r"N_{\mathrm{avg}}="
            r"\frac{N_1+N_2}{2}"
        )

        geiger_default = pd.DataFrame(
            {
                "Voltage V (V)": [
                    0, 60, 100, 160, 200,
                    260, 300, 360, 400, 460,
                    500, 560, 600, 660, 700,
                    760, 800, 860, 900, 960
                ],

                "N1": [
                    0, 0, 0, 0, 0,
                    0, 0, 0, 0, 0,
                    0, 942, 1241, 2456,
                    2645, 2798, 2819,
                    2957, 3043, 3158
                ],

                "N2": [
                    0, 0, 0, 0, 0,
                    0, 0, 0, 0, 0,
                    0, 927, 1296, 2424,
                    2707, 2745, 2882,
                    2970, 3090, 3126
                ]
            }
        )

        geiger_edit = st.data_editor(
            geiger_default,
            num_rows="dynamic",
            use_container_width=True,
            key="geiger_data"
        )

        geiger = safe_numeric_dataframe(
            geiger_edit,
            [
                "Voltage V (V)",
                "N1",
                "N2"
            ]
        )

        geiger = geiger.sort_values(
            "Voltage V (V)"
        ).copy()

        geiger[
            "N average"
        ] = (
            geiger["N1"]
            +
            geiger["N2"]
        ) / 2

        st.dataframe(
            geiger,
            use_container_width=True,
            hide_index=True
        )

        geiger_fig = None

        if len(geiger) >= 2:

            gv = geiger[
                "Voltage V (V)"
            ].to_numpy(float)

            gn = geiger[
                "N average"
            ].to_numpy(float)

            geiger_fig, ax = plt.subplots(
                figsize=(11, 5)
            )

            ax.plot(
                gv,
                gn,
                marker="o",
                label="Experimental Data"
            )

            ax.set_xlabel(
                "Voltage V (V)"
            )

            ax.set_ylabel(
                "Average Count Rate"
            )

            ax.set_title(
                "Geiger-Muller Operating Curve"
            )

            ax.grid(
                True,
                alpha=0.3
            )

            ax.legend()

            geiger_fig.tight_layout()

            st.pyplot(
                geiger_fig
            )

        st.header(
            "🧮 حساب منطقة الهضبة"
        )

        c1, c2, c3 = st.columns(3)

        with c1:

            vs = st.number_input(
                "Vs (V)",
                value=550.0,
                step=10.0,
                key="geiger_vs"
            )

        with c2:

            v1 = st.number_input(
                "V₁ (V)",
                value=650.0,
                step=10.0,
                key="geiger_v1"
            )

        with c3:

            v2 = st.number_input(
                "V₂ (V)",
                value=950.0,
                step=10.0,
                key="geiger_v2"
            )

        c4, c5 = st.columns(2)

        with c4:

            n1 = st.number_input(
                "N₁ عند V₁",
                value=2440.0,
                step=1.0,
                key="geiger_n1"
            )

        with c5:

            n2 = st.number_input(
                "N₂ عند V₂",
                value=3142.0,
                step=1.0,
                key="geiger_n2"
            )

        L = (
            v2 - v1
        )

        st.subheader(
            "1️⃣ طول منطقة الهضبة"
        )

        st.latex(
            r"L=V_2-V_1"
        )

        st.latex(
            rf"L="
            rf"{v2:.0f}"
            rf"-"
            rf"{v1:.0f}"
            rf"="
            rf"{L:.0f}"
            rf"\,\mathrm{{V}}"
        )

        st.subheader(
            "2️⃣ ميل منطقة الهضبة"
        )

        st.latex(
            r"S="
            r"\frac{N_2-N_1}{N_1}"
            r"\times"
            r"\frac{100}{V_2-V_1}"
            r"\times100"
        )

        slope = None

        if (
            n1 != 0
            and
            L != 0
        ):

            slope = (
                (
                    (n2 - n1)
                    /
                    n1
                )
                *
                (
                    100 / L
                )
                *
                100
            )

            st.latex(
                rf"S="
                rf"\frac{{"
                rf"{n2:.0f}-{n1:.0f}"
                rf"}}"
                rf"{{{n1:.0f}}}"
                rf"\times"
                rf"\frac{{100}}"
                rf"{{{v2:.0f}-{v1:.0f}}}"
                rf"\times100"
            )

            st.success(
                f"ميل الهضبة = "
                f"{slope:.4f} %/100 V"
            )

        operating_voltage = (
            v1 + v2
        ) / 2

        st.subheader(
            "3️⃣ جهد التشغيل"
        )

        st.latex(
            r"V_{\mathrm{operating}}="
            r"\frac{V_1+V_2}{2}"
        )

        st.latex(
            rf"V_{{\mathrm{{operating}}}}="
            rf"\frac{{"
            rf"{v1:.0f}+{v2:.0f}"
            rf"}}{{2}}"
            rf"="
            rf"{operating_voltage:.0f}"
            rf"\,\mathrm{{V}}"
        )

        st.metric(
            "جهد التشغيل المناسب",
            f"{operating_voltage:.0f} V"
        )

        # --------------------------------------------
        # R² لمنطقة الهضبة
        # --------------------------------------------

        st.subheader(
            "📊 Linear Regression و R² للهضبة"
        )

        plateau = geiger[
            (
                geiger["Voltage V (V)"] >= v1
            )
            &
            (
                geiger["Voltage V (V)"] <= v2
            )
        ].copy()

        regression = None
        regression_fig = None

        if len(plateau) >= 2:

            regression = calculate_regression(
                plateau[
                    "Voltage V (V)"
                ],
                plateau[
                    "N average"
                ]
            )

        if regression is not None:

            st.latex(
                rf"y="
                rf"({regression['slope']:.6g})x"
                rf"+"
                rf"({regression['intercept']:.6g})"
            )

            st.latex(
                rf"R^2="
                rf"{regression['r2']:.6f}"
            )

            regression_fig, ax_reg = plt.subplots(
                figsize=(10, 5)
            )

            plateau_x = plateau[
                "Voltage V (V)"
            ].to_numpy(float)

            plateau_y = plateau[
                "N average"
            ].to_numpy(float)

            predicted_y = (
                regression["slope"]
                *
                plateau_x
                +
                regression["intercept"]
            )

            order = np.argsort(
                plateau_x
            )

            ax_reg.scatter(
                plateau_x,
                plateau_y,
                label="Plateau Data"
            )

            ax_reg.plot(
                plateau_x[order],
                predicted_y[order],
                label="Linear Fit"
            )

            ax_reg.set_xlabel(
                "Voltage V (V)"
            )

            ax_reg.set_ylabel(
                "Average Count Rate"
            )

            ax_reg.set_title(
                "Geiger Plateau Linear Regression"
            )

            ax_reg.grid(
                True,
                alpha=0.3
            )

            ax_reg.legend()

            regression_fig.tight_layout()

            st.pyplot(
                regression_fig
            )

        else:

            st.info(
                "لا توجد نقاط كافية في منطقة الهضبة "
                "لحساب Linear Regression وR²."
            )

        # --------------------------------------------
        # تقرير جايجر
        # --------------------------------------------

        geiger_results = {
            "Vs": f"{vs:.6g} V",
            "V₁": f"{v1:.6g} V",
            "V₂": f"{v2:.6g} V",
            "طول الهضبة": f"{L:.6g} V",
            "جهد التشغيل": (
                f"{operating_voltage:.6g} V"
            ),
        }

        if slope is not None:

            geiger_results[
                "ميل الهضبة"
            ] = f"{slope:.6g} %/100 V"

        if regression is not None:

            geiger_results[
                "معادلة Linear Regression"
            ] = (
                f"y = "
                f"{regression['slope']:.6g}x + "
                f"{regression['intercept']:.6g}"
            )

            geiger_results[
                "R²"
            ] = f"{regression['r2']:.6f}"

        geiger_calculations = [
            (
                "حساب متوسط العد لكل قراءة: "
                "N_avg = (N₁ + N₂) / 2."
            ),

            (
                f"طول منطقة الهضبة: "
                f"L = V₂ - V₁ = "
                f"{v2:.6g} - {v1:.6g} "
                f"= {L:.6g} V."
            ),

            (
                f"جهد التشغيل: "
                f"V_operating = (V₁ + V₂)/2 = "
                f"({v1:.6g} + {v2:.6g})/2 "
                f"= {operating_voltage:.6g} V."
            )
        ]

        if slope is not None:

            geiger_calculations.append(
                f"ميل الهضبة: "
                f"S = [({n2:.6g}-{n1:.6g})/"
                f"{n1:.6g}] × "
                f"[100/({v2:.6g}-{v1:.6g})] × 100 "
                f"= {slope:.6g} %/100 V."
            )

        if regression is not None:

            geiger_calculations.extend(
                [
                    (
                        f"معادلة الانحدار الخطي للهضبة: "
                        f"y = "
                        f"{regression['slope']:.6g}x + "
                        f"{regression['intercept']:.6g}."
                    ),

                    (
                        f"معامل التحديد: "
                        f"R² = {regression['r2']:.6f}."
                    )
                ]
            )

        geiger_csv = geiger.to_csv(
            index=False
        ).encode(
            "utf-8-sig"
        )

        geiger_excel = create_excel_file(
            {
                "Readings": geiger,
                "Plateau": plateau
            },
            geiger_results
        )

        geiger_word = create_experiment_report(
            experiment_title=(
                "تحديد منطقة التشغيل لعداد جايجر"
            ),
            objective=(
                "رسم منحنى تشغيل عداد جايجر وتحديد "
                "منطقة الهضبة وجهد التشغيل المناسب."
            ),
            theory=(
                "يعتمد عداد جايجر-مولر على تسجيل "
                "الأحداث المؤينة. يتم اختيار جهد التشغيل "
                "ضمن منطقة الهضبة التي يتغير فيها معدل "
                "العد تغيرًا محدودًا مع الجهد."
            ),
            dataframe=geiger,
            formulas=[
                "N_avg = (N₁ + N₂) / 2",
                "L = V₂ - V₁",
                (
                    "S = [(N₂-N₁)/N₁] × "
                    "[100/(V₂-V₁)] × 100"
                ),
                "V_operating = (V₁ + V₂) / 2",
                "Linear fit: y = mx + b",
                "R² = coefficient of determination"
            ],
            calculations=geiger_calculations,
            results=geiger_results,
            figure=geiger_fig,
            analysis_text=(
                (
                    f"تم إجراء الانحدار الخطي على النقاط "
                    f"الواقعة بين V₁={v1:.6g} V "
                    f"وV₂={v2:.6g} V. "
                    f"بلغ R² = "
                    f"{regression['r2']:.6f}."
                )
                if regression is not None
                else
                "لم تتوفر نقاط كافية لحساب R²."
            )
        )

        st.divider()

        st.subheader(
            "📤 تصدير تجربة جايجر"
        )

        e1, e2, e3 = st.columns(3)

        with e1:

            st.download_button(
                "📄 CSV",
                geiger_csv,
                "PHIZKA_Geiger.csv",
                "text/csv"
            )

        with e2:

            st.download_button(
                "📊 Excel",
                geiger_excel,
                "PHIZKA_Geiger.xlsx",
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                )
            )

        with e3:

            st.download_button(
                "📝 تقرير Word كامل",
                geiger_word,
                "PHIZKA_Geiger_Report.docx",
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                )
            )

        if geiger_fig is not None:

            plt.close(
                geiger_fig
            )

        if regression_fig is not None:

            plt.close(
                regression_fig
            )


# ============================================================
# الحاسبة الفيزيائية
# ============================================================

elif page == "🧮 الحاسبة الفيزيائية":

    st.header(
        "🧮 الحاسبة الفيزيائية للطلبة"
    )

    calculator = st.selectbox(
        "اختاري نوع الحاسبة",
        [
            "حاسبة علمية",
            "قانون أوم",
            "القدرة الكهربائية",
            "الطاقة الحركية",
            "طاقة الفوتون",
            "الطول الموجي والتردد",
            "تحويل الوحدات",
        ]
    )

    if calculator == "حاسبة علمية":

        number = st.number_input(
            "العدد",
            value=1.0
        )

        operation = st.selectbox(
            "العملية",
            [
                "√x",
                "x²",
                "x³",
                "sin",
                "cos",
                "tan",
                "ln",
                "log10",
            ]
        )

        if operation == "√x":

            if number >= 0:

                result = math.sqrt(
                    number
                )

                st.latex(
                    rf"\sqrt{{{number}}}"
                    rf"={result}"
                )

            else:

                st.error(
                    "لا يمكن حساب الجذر الحقيقي لعدد سالب."
                )

        elif operation == "x²":

            st.latex(
                rf"({number})^2="
                rf"{number**2}"
            )

        elif operation == "x³":

            st.latex(
                rf"({number})^3="
                rf"{number**3}"
            )

        elif operation == "sin":

            result = math.sin(
                math.radians(number)
            )

            st.latex(
                rf"\sin({number}^\circ)="
                rf"{result:.8f}"
            )

        elif operation == "cos":

            result = math.cos(
                math.radians(number)
            )

            st.latex(
                rf"\cos({number}^\circ)="
                rf"{result:.8f}"
            )

        elif operation == "tan":

            result = math.tan(
                math.radians(number)
            )

            st.latex(
                rf"\tan({number}^\circ)="
                rf"{result:.8f}"
            )

        elif operation == "ln":

            if number > 0:

                result = math.log(
                    number
                )

                st.latex(
                    rf"\ln({number})="
                    rf"{result:.8f}"
                )

            else:

                st.error(
                    "ln يحتاج قيمة موجبة."
                )

        elif operation == "log10":

            if number > 0:

                result = math.log10(
                    number
                )

                st.latex(
                    rf"\log_{{10}}({number})="
                    rf"{result:.8f}"
                )

            else:

                st.error(
                    "log10 يحتاج قيمة موجبة."
                )


    elif calculator == "قانون أوم":

        st.latex(
            r"V=IR"
        )

        I_value = st.number_input(
            "I (A)",
            value=1.0,
            key="ohm_i"
        )

        R_value = st.number_input(
            "R (Ω)",
            value=1.0,
            key="ohm_r"
        )

        V_value = (
            I_value
            *
            R_value
        )

        st.latex(
            rf"V="
            rf"({I_value})"
            rf"({R_value})="
            rf"{V_value}"
            rf"\,\mathrm{{V}}"
        )


    elif calculator == "القدرة الكهربائية":

        st.latex(
            r"P=VI"
        )

        V_value = st.number_input(
            "V (V)",
            value=1.0,
            key="power_v"
        )

        I_value = st.number_input(
            "I (A)",
            value=1.0,
            key="power_i"
        )

        P_value = (
            V_value
            *
            I_value
        )

        st.latex(
            rf"P="
            rf"({V_value})"
            rf"({I_value})="
            rf"{P_value}"
            rf"\,\mathrm{{W}}"
        )


    elif calculator == "الطاقة الحركية":

        st.latex(
            r"K=\frac12mv^2"
        )

        mass = st.number_input(
            "m (kg)",
            min_value=0.0,
            value=1.0
        )

        velocity = st.number_input(
            "v (m/s)",
            value=1.0
        )

        kinetic_energy = (
            0.5
            *
            mass
            *
            velocity**2
        )

        st.latex(
            rf"K="
            rf"\frac12"
            rf"({mass})"
            rf"({velocity})^2="
            rf"{kinetic_energy}"
            rf"\,\mathrm{{J}}"
        )


    elif calculator == "طاقة الفوتون":

        st.latex(
            r"E=hf"
        )

        frequency = st.number_input(
            "f (Hz)",
            min_value=0.0,
            value=5e14,
            format="%.5e"
        )

        h = 6.62607015e-34

        energy = (
            h
            *
            frequency
        )

        st.latex(
            rf"E="
            rf"({h:.8e})"
            rf"({frequency:.5e})="
            rf"{energy:.8e}"
            rf"\,\mathrm{{J}}"
        )


    elif calculator == "الطول الموجي والتردد":

        st.latex(
            r"c=\lambda f"
        )

        frequency = st.number_input(
            "f (Hz)",
            min_value=0.000001,
            value=5e14,
            format="%.5e",
            key="wave_frequency"
        )

        c = 299792458.0

        wavelength_value = (
            c
            /
            frequency
        )

        st.latex(
            rf"\lambda="
            rf"\frac{{{c:.0f}}}"
            rf"{{{frequency:.5e}}}="
            rf"{wavelength_value:.8e}"
            rf"\,\mathrm{{m}}"
        )


    elif calculator == "تحويل الوحدات":

        conversion = st.selectbox(
            "التحويل",
            [
                "m → cm",
                "cm → m",
                "m → nm",
                "nm → m",
                "A → mA",
                "mA → A",
                "W → mW",
                "mW → W",
                "eV → J",
                "J → eV",
            ]
        )

        value = st.number_input(
            "القيمة",
            value=1.0
        )

        electron_charge = (
            1.602176634e-19
        )

        if conversion == "m → cm":

            result = value * 100
            unit = "cm"

        elif conversion == "cm → m":

            result = value / 100
            unit = "m"

        elif conversion == "m → nm":

            result = value * 1e9
            unit = "nm"

        elif conversion == "nm → m":

            result = value * 1e-9
            unit = "m"

        elif conversion == "A → mA":

            result = value * 1000
            unit = "mA"

        elif conversion == "mA → A":

            result = value / 1000
            unit = "A"

        elif conversion == "W → mW":

            result = value * 1000
            unit = "mW"

        elif conversion == "mW → W":

            result = value / 1000
            unit = "W"

        elif conversion == "eV → J":

            result = (
                value
                *
                electron_charge
            )

            unit = "J"

        else:

            result = (
                value
                /
                electron_charge
            )

            unit = "eV"

        st.success(
            f"{value} = "
            f"{result:.10g} {unit}"
        )


# ============================================================
# تحليل البيانات
# ============================================================

elif page == "📊 تحليل البيانات":

    st.header(
        "📊 تحليل البيانات"
    )

    st.write(
        "ارفعي ملف CSV أو Excel لإجراء الرسم "
        "والانحدار الخطي وحساب R²."
    )

    uploaded_file = st.file_uploader(
        "📂 رفع CSV أو Excel",
        type=[
            "csv",
            "xlsx",
            "xls"
        ],
        key="analysis_upload"
    )

    if uploaded_file is not None:

        try:

            if uploaded_file.name.lower().endswith(
                ".csv"
            ):

                uploaded_df = pd.read_csv(
                    uploaded_file
                )

            else:

                uploaded_df = pd.read_excel(
                    uploaded_file
                )

            st.dataframe(
                uploaded_df,
                use_container_width=True
            )

            if len(
                uploaded_df.columns
            ) >= 2:

                c1, c2 = st.columns(2)

                with c1:

                    x_column = st.selectbox(
                        "X",
                        uploaded_df.columns,
                        key="analysis_x"
                    )

                with c2:

                    y_options = [
                        col
                        for col
                        in uploaded_df.columns
                        if col != x_column
                    ]

                    y_column = st.selectbox(
                        "Y",
                        y_options,
                        key="analysis_y"
                    )

                analysis_df = safe_numeric_dataframe(
                    uploaded_df,
                    [
                        x_column,
                        y_column
                    ]
                )

                if len(analysis_df) >= 2:

                    regression = calculate_regression(
                        analysis_df[
                            x_column
                        ],
                        analysis_df[
                            y_column
                        ]
                    )

                    fig, ax = plt.subplots(
                        figsize=(10, 5)
                    )

                    ax.scatter(
                        analysis_df[
                            x_column
                        ],
                        analysis_df[
                            y_column
                        ],
                        label="Data"
                    )

                    if regression is not None:

                        x_array = analysis_df[
                            x_column
                        ].to_numpy(float)

                        y_fit = (
                            regression["slope"]
                            *
                            x_array
                            +
                            regression[
                                "intercept"
                            ]
                        )

                        order = np.argsort(
                            x_array
                        )

                        ax.plot(
                            x_array[order],
                            y_fit[order],
                            label="Linear Fit"
                        )

                        st.latex(
                            rf"y="
                            rf"({regression['slope']:.6g})x"
                            rf"+"
                            rf"({regression['intercept']:.6g})"
                        )

                        st.latex(
                            rf"R^2="
                            rf"{regression['r2']:.6f}"
                        )

                    ax.set_xlabel(
                        str(x_column)
                    )

                    ax.set_ylabel(
                        str(y_column)
                    )

                    ax.grid(
                        True,
                        alpha=0.3
                    )

                    ax.legend()

                    fig.tight_layout()

                    st.pyplot(
                        fig
                    )

                    plt.close(
                        fig
                    )

        except Exception as error:

            st.error(
                f"تعذر قراءة الملف: {error}"
            )


# ============================================================
# GNUPLOT 3D
# ============================================================

elif page == "🧊 Gnuplot 3D":

    st.header(
        "🧊 Gnuplot 3D"
    )

    st.write(
        "عدلي قيم X وY وZ ثم اضغطي "
        "«إنشاء الرسم 3D»."
    )

    default_3d = pd.DataFrame(
        {
            "X": [
                -2, -1, 0, 1, 2,
                -2, -1, 0, 1, 2,
                -2, -1, 0, 1, 2,
                -2, -1, 0, 1, 2,
                -2, -1, 0, 1, 2
            ],

            "Y": [
                -2, -2, -2, -2, -2,
                -1, -1, -1, -1, -1,
                0, 0, 0, 0, 0,
                1, 1, 1, 1, 1,
                2, 2, 2, 2, 2
            ],

            "Z": [
                8, 5, 4, 5, 8,
                5, 2, 1, 2, 5,
                4, 1, 0, 1, 4,
                5, 2, 1, 2, 5,
                8, 5, 4, 5, 8
            ]
        }
    )

    data_3d_edit = st.data_editor(
        default_3d,
        num_rows="dynamic",
        use_container_width=True,
        key="gnuplot_3d_data"
    )

    data_3d = safe_numeric_dataframe(
        data_3d_edit,
        [
            "X",
            "Y",
            "Z"
        ]
    )

    c1, c2 = st.columns(2)

    with c1:

        title_3d = st.text_input(
            "عنوان الرسم",
            value="PHIZKA 3D Surface"
        )

        plot_mode = st.selectbox(
            "نوع الرسم",
            [
                "Surface",
                "Points",
                "Lines + Points",
                "Color-mapped Surface"
            ]
        )

    with c2:

        view_angle = st.slider(
            "زاوية العرض",
            min_value=10,
            max_value=85,
            value=60,
            step=5
        )

        x_label_3d = st.text_input(
            "اسم محور X",
            value="X"
        )

        y_label_3d = st.text_input(
            "اسم محور Y",
            value="Y"
        )

        z_label_3d = st.text_input(
            "اسم محور Z",
            value="Z"
        )

    if gnuplot_path is None:

        st.error(
            "Gnuplot غير مكتشف على الجهاز."
        )

    else:

        st.success(
            "Gnuplot جاهز ✅"
        )

        if st.button(
            "🧊 إنشاء الرسم 3D",
            type="primary"
        ):

            if len(data_3d) < 3:

                st.warning(
                    "أدخلي ثلاث نقاط على الأقل."
                )

            else:

                try:

                    temp_dir = tempfile.gettempdir()

                    data_file = os.path.join(
                        temp_dir,
                        "phizka_3d.dat"
                    )

                    script_file = os.path.join(
                        temp_dir,
                        "phizka_3d.gp"
                    )

                    output_file = os.path.join(
                        temp_dir,
                        "phizka_3d.png"
                    )

                    sorted_data = data_3d.sort_values(
                        [
                            "Y",
                            "X"
                        ]
                    )

                    with open(
                        data_file,
                        "w",
                        encoding="utf-8"
                    ) as file:

                        previous_y = None

                        for row in sorted_data.itertuples(
                            index=False,
                            name=None
                        ):

                            x_value, y_value, z_value = row

                            if (
                                previous_y is not None
                                and
                                y_value != previous_y
                            ):

                                file.write(
                                    "\n"
                                )

                            file.write(
                                f"{x_value} "
                                f"{y_value} "
                                f"{z_value}\n"
                            )

                            previous_y = y_value

                    data_path = data_file.replace(
                        "\\",
                        "/"
                    )

                    output_path = output_file.replace(
                        "\\",
                        "/"
                    )

                    if plot_mode == "Surface":

                        plot_command = (
                            f'splot "{data_path}" '
                            f'using 1:2:3 '
                            f'with lines '
                            f'title "Surface"'
                        )

                    elif plot_mode == "Points":

                        plot_command = (
                            f'splot "{data_path}" '
                            f'using 1:2:3 '
                            f'with points '
                            f'pointtype 7 '
                            f'pointsize 1.4 '
                            f'title "Points"'
                        )

                    elif plot_mode == "Lines + Points":

                        plot_command = (
                            f'splot "{data_path}" '
                            f'using 1:2:3 '
                            f'with linespoints '
                            f'title "Data"'
                        )

                    else:

                        plot_command = (
                            f'splot "{data_path}" '
                            f'using 1:2:3 '
                            f'with pm3d '
                            f'title "Surface"'
                        )

                    script = f"""
set terminal pngcairo size 1200,800
set output "{output_path}"
set title "{title_3d}"
set xlabel "{x_label_3d}"
set ylabel "{y_label_3d}"
set zlabel "{z_label_3d}"
set grid
set hidden3d
set view {view_angle},35
set palette rgbformulae 33,13,10
{plot_command}
"""

                    with open(
                        script_file,
                        "w",
                        encoding="utf-8"
                    ) as file:

                        file.write(
                            script
                        )

                    subprocess.run(
                        [
                            gnuplot_path,
                            script_file
                        ],
                        check=True
                    )

                    if os.path.exists(
                        output_file
                    ):

                        st.image(
                            output_file,
                            caption=title_3d
                        )

                        with open(
                            output_file,
                            "rb"
                        ) as image_file:

                            image_bytes = (
                                image_file.read()
                            )

                        st.download_button(
                            "📥 تنزيل الرسم 3D",
                            image_bytes,
                            "PHIZKA_3D.png",
                            "image/png"
                        )

                    else:

                        st.error(
                            "لم يتم إنشاء الرسم."
                        )

                except Exception as error:

                    st.error(
                        f"حدث خطأ في Gnuplot: {error}"
                    )


# ============================================================
# ورشة تحليل البيانات الفيزيائية
# ============================================================

elif page == "🧪 ورشة تحليل البيانات الفيزيائية":

    st.header("🧪 ورشة تحليل البيانات الفيزيائية")

    st.write(
        "ورشة تطبيقية داخل PHIZKA لتحويل القراءات التجريبية إلى "
        "جدول منظم، رسم بياني، انحدار خطي، معادلة أفضل خط مستقيم، "
        "ومعامل التحديد R²، ثم تصدير النتائج إلى CSV وExcel وWord."
    )

    workshop_experiment = st.selectbox(
        "اختاري التطبيق العملي",
        ["🪝 قانون هوك | Hooke's Law"],
        key="workshop_experiment"
    )

    if workshop_experiment == "🪝 قانون هوك | Hooke's Law":

        st.subheader("🪝 تحليل بيانات قانون هوك")

        st.markdown(
            "**الهدف:** دراسة العلاقة بين القوة المؤثرة والاستطالة، "
            "واستخراج ثابت النابض من ميل الرسم البياني."
        )

        st.latex(r"F=kx")
        st.latex(r"F=mx+b")
        st.write("في الرسم F مقابل x يكون ميل الخط m هو ثابت النابض k بوحدة N/m.")

        hooke_default = pd.DataFrame(
            {
                "Extension x (m)": [0.01, 0.02, 0.03, 0.04, 0.05, 0.06],
                "Force F (N)": [0.50, 1.01, 1.49, 2.02, 2.51, 3.00],
            }
        )

        st.subheader("📋 القراءات التجريبية")
        st.caption("القيم التالية مثال جاهز للورشة ويمكن تعديلها بالكامل إلى قراءات التجربة الحقيقية.")

        hooke_edit = st.data_editor(
            hooke_default,
            num_rows="dynamic",
            use_container_width=True,
            key="hooke_workshop_table"
        )

        hooke = safe_numeric_dataframe(
            hooke_edit,
            ["Extension x (m)", "Force F (N)"]
        ).sort_values("Extension x (m)").copy()

        hooke_regression = None
        hooke_fig = None

        if len(hooke) >= 2:
            hooke_regression = calculate_regression(
                hooke["Extension x (m)"],
                hooke["Force F (N)"]
            )

        if hooke_regression is not None:
            x_values = hooke["Extension x (m)"].to_numpy(float)
            f_values = hooke["Force F (N)"].to_numpy(float)
            predicted = (
                hooke_regression["slope"] * x_values
                + hooke_regression["intercept"]
            )

            hooke["F predicted (N)"] = predicted
            hooke["Residual (N)"] = f_values - predicted

            st.subheader("📊 النتائج")
            c1, c2, c3 = st.columns(3)
            c1.metric("ثابت النابض k", f"{hooke_regression['slope']:.4f} N/m")
            c2.metric("المقطع b", f"{hooke_regression['intercept']:.4f} N")
            c3.metric("R²", f"{hooke_regression['r2']:.6f}")

            st.latex(
                rf"F=({hooke_regression['slope']:.6g})x"
                rf"+({hooke_regression['intercept']:.6g})"
            )
            st.latex(rf"R^2={hooke_regression['r2']:.6f}")

            hooke_fig, ax = plt.subplots(figsize=(10, 5.5))
            ax.scatter(x_values, f_values, s=70, label="Experimental Data")

            order = np.argsort(x_values)
            ax.plot(
                x_values[order],
                predicted[order],
                linewidth=2,
                label="Linear Fit"
            )
            ax.set_xlabel("Extension x (m)")
            ax.set_ylabel("Force F (N)")
            ax.set_title("Hooke's Law: Force vs Extension")
            ax.grid(True, alpha=0.3)
            ax.legend()
            hooke_fig.tight_layout()
            st.pyplot(hooke_fig)

            st.subheader("🔎 تفسير R²")
            r2 = hooke_regression["r2"]
            if r2 >= 0.99:
                st.success("العلاقة خطية بدرجة عالية جدًا، والبيانات متوافقة بقوة مع نموذج قانون هوك ضمن مدى القياس.")
            elif r2 >= 0.95:
                st.info("العلاقة الخطية قوية، مع وجود انحرافات تجريبية بسيطة عن الخط المستقيم.")
            else:
                st.warning("البيانات تُظهر انحرافًا ملحوظًا عن الخطية؛ راجعي القراءات أو مدى مرونة النابض أو عدم اليقين التجريبي.")

            st.subheader("🧮 التعويض والحساب")
            hooke_calculations = [
                "نستخدم النموذج الخطي: F = mx + b.",
                f"من الانحدار الخطي: m = {hooke_regression['slope']:.6g} N/m.",
                f"إذن ثابت النابض: k = m = {hooke_regression['slope']:.6g} N/m.",
                f"المقطع: b = {hooke_regression['intercept']:.6g} N.",
                f"معامل التحديد: R² = {hooke_regression['r2']:.6f}.",
            ]

            for line in hooke_calculations:
                st.write("• " + line)

            st.subheader("🧠 توقع قوة من الاستطالة")
            x_predict = st.number_input(
                "أدخلي استطالة x (m)",
                min_value=0.0,
                value=0.035,
                step=0.005,
                format="%.4f",
                key="hooke_x_predict"
            )
            f_predict = (
                hooke_regression["slope"] * x_predict
                + hooke_regression["intercept"]
            )
            st.latex(
                rf"F=({hooke_regression['slope']:.6g})"
                rf"({x_predict:.6g})+({hooke_regression['intercept']:.6g})"
                rf"={f_predict:.6g}\,\mathrm{{N}}"
            )

            st.subheader("📋 جدول التحليل")
            st.dataframe(hooke, use_container_width=True, hide_index=True)

            hooke_results = {
                "ثابت النابض k": f"{hooke_regression['slope']:.6g} N/m",
                "المقطع b": f"{hooke_regression['intercept']:.6g} N",
                "R²": f"{hooke_regression['r2']:.6f}",
                "عدد القراءات": len(hooke),
            }

            hooke_csv = hooke.to_csv(index=False).encode("utf-8-sig")
            hooke_excel = create_excel_file(
                {"Hooke Data": hooke},
                hooke_results
            )
            hooke_word = create_experiment_report(
                experiment_title="ورشة تحليل البيانات الفيزيائية - قانون هوك",
                objective=(
                    "دراسة العلاقة بين القوة والاستطالة وتطبيق الانحدار "
                    "الخطي لاستخراج ثابت النابض وتقييم جودة المطابقة باستخدام R²."
                ),
                theory=(
                    "ينص قانون هوك، ضمن حد المرونة، على أن القوة المؤثرة "
                    "تتناسب طرديًا مع استطالة النابض: F = kx. عند تمثيل F "
                    "مقابل x يكون ميل أفضل خط مستقيم مساويًا لثابت النابض k."
                ),
                dataframe=hooke,
                formulas=[
                    "F = kx",
                    "F = mx + b",
                    "k = slope = m",
                    "R² = coefficient of determination"
                ],
                calculations=hooke_calculations,
                results=hooke_results,
                figure=hooke_fig,
                analysis_text=(
                    f"أعطى الانحدار الخطي R² = {hooke_regression['r2']:.6f}. "
                    "كلما اقتربت R² من 1 دل ذلك على أن النموذج الخطي يصف "
                    "البيانات التجريبية بدرجة أفضل."
                )
            )

            st.divider()
            st.subheader("📤 تصدير نتائج الورشة")
            e1, e2, e3 = st.columns(3)
            with e1:
                st.download_button(
                    "📄 CSV",
                    hooke_csv,
                    "PHIZKA_Hooke_Workshop.csv",
                    "text/csv"
                )
            with e2:
                st.download_button(
                    "📊 Excel",
                    hooke_excel,
                    "PHIZKA_Hooke_Workshop.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            with e3:
                st.download_button(
                    "📝 تقرير Word كامل",
                    hooke_word,
                    "PHIZKA_Hooke_Workshop_Report.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            if hooke_fig is not None:
                plt.close(hooke_fig)

        else:
            st.info("أدخلي قراءتين صالحـتين على الأقل حتى يتم إجراء الانحدار الخطي وحساب R².")


# ============================================================
# مشروع التخرج والبحث العلمي
# ============================================================

elif page == "🎓 مشروع التخرج والبحث العلمي":

    st.header(
        "🎓 مشروع التخرج والبحث العلمي"
    )

    project_name = st.text_input(
        "اسم المشروع"
    )

    project_goal = st.text_area(
        "هدف المشروع"
    )

    project_upload = st.file_uploader(
        "📂 رفع CSV أو Excel",
        type=[
            "csv",
            "xlsx",
            "xls"
        ],
        key="project_upload"
    )

    if project_upload is not None:

        try:

            if project_upload.name.lower().endswith(
                ".csv"
            ):

                project_default = pd.read_csv(
                    project_upload
                )

            else:

                project_default = pd.read_excel(
                    project_upload
                )

        except Exception as error:

            st.error(
                f"تعذر قراءة الملف: {error}"
            )

            project_default = pd.DataFrame(
                {
                    "X": [0.0, 1.0, 2.0],
                    "Y": [0.0, 1.0, 2.0]
                }
            )

    else:

        project_default = pd.DataFrame(
            {
                "X": [
                    0.0,
                    1.0,
                    2.0
                ],

                "Y": [
                    0.0,
                    1.0,
                    2.0
                ]
            }
        )

    project_data = st.data_editor(
        project_default,
        num_rows="dynamic",
        use_container_width=True,
        key="project_data"
    )

    if len(
        project_data.columns
    ) >= 2:

        c1, c2 = st.columns(2)

        with c1:

            project_x = st.selectbox(
                "محور X",
                project_data.columns,
                key="project_x"
            )

        with c2:

            y_options = [
                col
                for col
                in project_data.columns
                if col != project_x
            ]

            project_y = st.selectbox(
                "محور Y",
                y_options,
                key="project_y"
            )

        project_clean = safe_numeric_dataframe(
            project_data,
            [
                project_x,
                project_y
            ]
        )

        if len(project_clean) >= 2:

            project_fig, ax = plt.subplots(
                figsize=(10, 5)
            )

            ax.plot(
                project_clean[
                    project_x
                ],
                project_clean[
                    project_y
                ],
                marker="o"
            )

            ax.set_xlabel(
                str(project_x)
            )

            ax.set_ylabel(
                str(project_y)
            )

            ax.set_title(
                project_name
                if project_name
                else
                "Scientific Project"
            )

            ax.grid(
                True,
                alpha=0.3
            )

            project_fig.tight_layout()

            st.pyplot(
                project_fig
            )

            if st.checkbox(
                "📊 حساب Linear Regression وR²"
            ):

                project_regression = calculate_regression(
                    project_clean[
                        project_x
                    ],
                    project_clean[
                        project_y
                    ]
                )

                if project_regression is not None:

                    st.latex(
                        rf"y="
                        rf"({project_regression['slope']:.6g})x"
                        rf"+"
                        rf"({project_regression['intercept']:.6g})"
                    )

                    st.latex(
                        rf"R^2="
                        rf"{project_regression['r2']:.6f}"
                    )

            plt.close(
                project_fig
            )

    st.divider()

    st.subheader(
        "🔜 التطوير القادم"
    )

    st.write(
        """
        • Curve Fitting غير الخطي  
        • Error Bars  
        • Uncertainty Analysis  
        • Anomaly Detection  
        • مقارنة عدة مجموعات بيانات  
        • أدوات الذكاء الاصطناعي للتحليل العلمي
        """
    )


# ============================================================
# نهاية PHIZKA
# ============================================================

st.divider()

st.caption(
    "⚛️ PHIZKA Scientific Platform | فيزكا العلمية"
)