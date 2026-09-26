import io
import math
from datetime import datetime

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from scipy.optimize import curve_fit
from scipy.stats import linregress
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

st.set_page_config(page_title='PHIZKA | ورشة تحليل البيانات الفيزيائية', page_icon='⚛️', layout='wide')
st.markdown('''<style>.block-container{padding-top:1.5rem;max-width:1250px}h1,h2,h3{text-align:right}div[data-testid="stMetric"]{border:1px solid #b8d8e8;border-radius:12px;padding:12px}</style>''', unsafe_allow_html=True)
st.title('⚛️ PHIZKA | ورشة تحليل البيانات الفيزيائية')
st.caption('من قراءات المختبر إلى الرسم والتحليل والتقرير — لأي تجربة فيزيائية')

HOOKE = pd.DataFrame({
    'الكتلة (g)': pd.Series([None]*5, dtype='float64'),
    'القوة (N)': pd.Series([None]*5, dtype='float64'),
    'الاستطالة (cm)': pd.Series([None]*5, dtype='float64')
})
DEMO = pd.DataFrame({'X':[0.,1.,2.,3.,4.], 'Y':[0.,2.1,3.9,6.2,8.1], 'Z':[0.,1.,4.,9.,16.]})

@st.cache_data
def example_3d():
    x,y=np.meshgrid(np.linspace(-2,2,11),np.linspace(-2,2,11))
    return pd.DataFrame({'X':x.ravel(),'Y':y.ravel(),'Z':(np.sin(x)*np.cos(y)).ravel()})

def clean(df, cols):
    out=df.copy()
    for c in cols: out[c]=pd.to_numeric(out[c],errors='coerce')
    return out.dropna(subset=cols).reset_index(drop=True)

def fit_line(x,y):
    x=np.asarray(x,dtype=float); y=np.asarray(y,dtype=float)
    if len(x)<2 or np.ptp(x)<1e-14: return None
    r=linregress(x,y)
    yhat=r.slope*x+r.intercept
    ss_tot=float(np.sum((y-y.mean())**2)); ss_res=float(np.sum((y-yhat)**2))
    r2=1-ss_res/ss_tot if ss_tot>1e-15 else (1. if ss_res<1e-15 else float('nan'))
    return {'slope':float(r.slope),'intercept':float(r.intercept),'r2':r2,'yhat':yhat}

def png(fig):
    b=io.BytesIO();fig.savefig(b,format='png',dpi=180,bbox_inches='tight');return b.getvalue()

def excel(df, results):
    b=io.BytesIO()
    with pd.ExcelWriter(b,engine='openpyxl') as w:
        df.to_excel(w,sheet_name='Readings',index=False)
        pd.DataFrame([{'المؤشر':k,'القيمة':v} for k,v in results.items()]).to_excel(w,sheet_name='Results',index=False)
    return b.getvalue()

def rtl(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    pp=p._p.get_or_add_pPr(); bidi=OxmlElement('w:bidi');bidi.set(qn('w:val'),'1');pp.append(bidi)

def report(title,objective,df,results,steps,graph,analysis):
    d=Document(); normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    p=d.add_paragraph('PHIZKA | فيزكا العلمية');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(title);rtl(p);p.runs[0].bold=True;p.runs[0].font.size=Pt(17)
    sections=[('معلومات التجربة',f'تاريخ التقرير: {datetime.now():%Y-%m-%d %H:%M}'),('الهدف',objective),('المنهج', 'أُدخلت القراءات في منصة PHIZKA ثم أُجري التحليل العددي والرسم البياني حسب اختيار الطالبة.'),('القوانين والتعويض', '\n'.join(steps)),('النتائج','\n'.join(f'{k}: {v}' for k,v in results.items())),('التحليل',analysis),('الخلاصة','النتائج محسوبة من القراءات الحالية، ويجب مراجعة الوحدات والقانون الفيزيائي قبل اعتماد الاستنتاج.')]
    for name,content in sections[:3]:
        p=d.add_paragraph(name,style='Heading 2');rtl(p)
        for line in content.split('\n'): p=d.add_paragraph(line);rtl(p)
    p=d.add_paragraph('جدول القراءات',style='Heading 2');rtl(p)
    table=d.add_table(rows=1,cols=len(df.columns));table.style='Table Grid'
    for j,c in enumerate(df.columns):table.cell(0,j).text=str(c)
    for row in df.itertuples(index=False,name=None):
        cells=table.add_row().cells
        for j,v in enumerate(row):cells[j].text=f'{v:.7g}' if isinstance(v,(float,np.floating)) else str(v)
    for name,content in sections[3:]:
        p=d.add_paragraph(name,style='Heading 2');rtl(p)
        for line in content.split('\n'):p=d.add_paragraph(line);rtl(p)
    if graph:
        p=d.add_paragraph('الرسم البياني',style='Heading 2');rtl(p)
        d.add_picture(io.BytesIO(graph),width=Inches(6.0))
    b=io.BytesIO();d.save(b);return b.getvalue()

def downloads(df,results,steps,graph,title,objective,analysis,prefix):
    st.subheader('📥 تصدير النتائج')
    c1,c2,c3,c4=st.columns(4)
    c1.download_button('CSV',df.to_csv(index=False).encode('utf-8-sig'),f'{prefix}.csv','text/csv')
    c2.download_button('Excel',excel(df,results),f'{prefix}.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    c3.download_button('Word',report(title,objective,df,results,steps,graph,analysis),f'{prefix}_report.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    if graph:c4.download_button('PNG',graph,f'{prefix}_plot.png','image/png')

st.sidebar.header('أقسام الورشة')
page=st.sidebar.radio('اختاري', ['🏠 البداية','🧪 قانون هوك — تجربة الدكتورة','📊 تجربة جديدة / ملف خارجي','🧊 مختبر 3D','📘 دليل الطالبة'])

if page=='🏠 البداية':
    st.info('يمكن للطالبة البدء بقانون هوك، أو إدخال قراءات أي تجربة أخرى، أو رفع CSV / Excel، ثم تحديد المحاور وحساب R² وتصدير التقرير.')
    a,b,c=st.columns(3)
    a.metric('تجربة الورشة','قانون هوك');b.metric('التحليل','Linear fit + R²');c.metric('الرسم','2D + 3D')
    st.markdown('**مسار العمل:** إدخال القراءات ← اختيار X وY ← تحديد النموذج ← قراءة الميل وR² ← إنشاء الرسم ← تنزيل التقرير.')
    st.warning('R² يصف ملاءمة النموذج للبيانات ولا يثبت صحة القانون الفيزيائي. الرسم 3D يحتاج ثلاثة متغيرات عددية، والسطح يحتاج شبكة بيانات مناسبة.')

elif page=='🧪 قانون هوك — تجربة الدكتورة':
    st.header('تجربة قانون هوك | الاستطالة وثابت الزنبرك')
    st.write('أدخلي قراءات التجربة الفعلية: الكتلة (g)، القوة (N)، والاستطالة (cm). الجدول يبدأ فارغًا ويمكنك إضافة صفوف أو حذفها، والرسم و k و R² تتحدث تلقائيًا حسب القراءات.')
    edited=st.data_editor(HOOKE,num_rows='dynamic',width='stretch',key='hooke_editor')
    df=clean(edited,list(HOOKE.columns));df=df[df['الاستطالة (cm)']>=0].copy()
    if len(df)<2:st.warning('أدخلي قراءتين صحيحتين على الأقل.');st.stop()
    df['الاستطالة (m)']=df['الاستطالة (cm)']/100
    fit=fit_line(df['الاستطالة (m)'],df['القوة (N)'])
    if fit is None:st.error('الاستطالات متطابقة؛ لا يمكن حساب الميل.');st.stop()
    k=fit['slope'];b=fit['intercept'];r2=fit['r2']
    c1,c2,c3=st.columns(3);c1.metric('ثابت الزنبرك k',f'{k:.4f} N/m');c2.metric('التقاطع b',f'{b:.4f} N');c3.metric('R²',f'{r2:.6f}')
    st.latex(r'F=k\,\Delta x+b')
    st.write(f'**معادلة الملاءمة:** F = ({k:.5f} N/m) × Δx + ({b:.5f} N)')
    fig,ax=plt.subplots(figsize=(9,5));x=df['الاستطالة (m)'].to_numpy();y=df['القوة (N)'].to_numpy()
    ax.scatter(x,y,s=65,label='Measured readings');xx=np.linspace(min(x),max(x),100);ax.plot(xx,k*xx+b,label=f'Fit: k={k:.3f} N/m, R²={r2:.5f}')
    ax.set(xlabel='Extension Δx (m)',ylabel='Force F (N)',title="Hooke's law");ax.grid(alpha=.3);ax.legend();st.pyplot(fig);graph=png(fig);plt.close(fig)
    st.dataframe(df,width='stretch',hide_index=True)
    steps=['F = k Δx + b','Δx(m) = Δx(cm) / 100','k = slope of F versus Δx in metres']+[f'قراءة {i+1}: Δx={row["الاستطالة (cm)"]:.4g}/100={row["الاستطالة (m)"]:.5g} m, F={row["القوة (N)"]:.4g} N' for i,(_,row) in enumerate(df.iterrows())]+[f'F = {k:.6g} Δx + {b:.6g}',f'R² = {r2:.7g}']
    downloads(df,{'k (N/m)':k,'intercept (N)':b,'R²':r2},steps,graph,'تجربة قانون هوك','حساب ثابت الزنبرك من ميل القوة مقابل الاستطالة.', 'تمت ملاءمة البيانات بخط مستقيم مع تقاطع حر. وحدة الاستطالة المستخدمة في الملاءمة هي المتر.','PHIZKA_Hooke')

elif page=='📊 تجربة جديدة / ملف خارجي':
    st.header('تحليل تجربة فيزيائية غير مضافة مسبقًا')
    st.caption('تُحدد الطالبة القانون الفيزيائي والمحاور والوحدات؛ المنصة لا تفترض قانونًا من البيانات وحدها.')
    name=st.text_input('اسم التجربة',value='تجربتي الفيزيائية')
    objective=st.text_area('هدف التجربة / القانون النظري',value='دراسة العلاقة بين المتغير المستقل والمتغير التابع.')
    source=st.radio('مصدر البيانات',['إدخال يدوي','CSV / Excel'],horizontal=True)
    if source=='CSV / Excel':
        upload=st.file_uploader('اختاري الملف',type=['csv','xlsx'])
        if upload is None:st.info('ارفعي ملفًا للمتابعة.');st.stop()
        try: raw=pd.read_csv(upload) if upload.name.lower().endswith('.csv') else pd.read_excel(upload)
        except Exception as e:st.error(f'تعذر قراءة الملف: {e}');st.stop()
    else:raw=DEMO[['X','Y']].copy()
    edited=st.data_editor(raw,num_rows='dynamic',width='stretch',key='custom_editor')
    cols=list(edited.columns)
    if len(cols)<2:st.warning('يلزم عمودان على الأقل.');st.stop()
    c1,c2=st.columns(2)
    xcol=c1.selectbox('المحور الأفقي X',cols,index=0)
    ycol=c2.selectbox('المحور الرأسي Y',[c for c in cols if c!=xcol])
    xlabel=st.text_input('وحدة X',value='');ylabel=st.text_input('وحدة Y',value='')
    df=clean(edited,[xcol,ycol]);st.caption(f'عدد القراءات الصالحة: {len(df)}')
    if len(df)<2:st.warning('أدخلي نقطتين صحيحتين على الأقل.');st.stop()
    x=df[xcol].to_numpy(float);y=df[ycol].to_numpy(float)
    mode=st.selectbox('نموذج التحليل',['خطي y = mx + b','تربيعي y = ax² + bx + c','رسم فقط'])
    fig,ax=plt.subplots(figsize=(9,5));ax.scatter(x,y,label='Experimental data',s=50)
    steps=[];results={};analysis='تم رسم البيانات دون ملاءمة رياضية.'
    if mode=='خطي y = mx + b':
        f=fit_line(x,y)
        if f:
            xx=np.linspace(min(x),max(x),200);ax.plot(xx,f['slope']*xx+f['intercept'],label='Linear fit')
            results={'الميل m':f['slope'],'التقاطع b':f['intercept'],'R²':f['r2']}
            steps=[f'y = ({f["slope"]:.7g})x + ({f["intercept"]:.7g})',f'R² = {f["r2"]:.7g}']
            analysis='R² معامل تحديد الملاءمة الخطية للقراءات الحالية، وليس برهانًا على صحة النموذج الفيزيائي.'
        else:st.warning('لا يمكن ملاءمة خطية لأن قيم X متطابقة.')
    elif mode=='تربيعي y = ax² + bx + c':
        if len(df)>=3 and len(np.unique(x))>=3:
            co=np.polyfit(x,y,2);pred=np.polyval(co,x);den=np.sum((y-y.mean())**2);r2=1-np.sum((y-pred)**2)/den if den>1e-15 else float('nan')
            xx=np.linspace(min(x),max(x),200);ax.plot(xx,np.polyval(co,xx),label='Quadratic fit')
            results={'a':co[0],'b':co[1],'c':co[2],'R²':r2};steps=[f'y = {co[0]:.7g}x² + {co[1]:.7g}x + {co[2]:.7g}',f'R² = {r2:.7g}']
            analysis='ملاءمة تربيعية؛ يجب التأكد أن هذا النموذج مناسب لنظرية التجربة.'
        else:st.warning('الملاءمة التربيعية تتطلب 3 قيم X مختلفة على الأقل.')
    ax.set(xlabel=f'{xcol} {xlabel}',ylabel=f'{ycol} {ylabel}',title=name);ax.grid(alpha=.3);ax.legend();st.pyplot(fig);graph=png(fig);plt.close(fig)
    for k,v in results.items():st.metric(k,f'{v:.6g}' if np.isfinite(v) else 'غير معرف')
    downloads(df,results,steps,graph,name,objective,analysis,'PHIZKA_Experiment')

elif page=='🧊 مختبر 3D':
    st.header('رسم ثلاثي الأبعاد لأي تجربة')
    st.write('يمكن رسم X وY وZ كنقاط أو سطح. السطح لا يُنشأ إلا عند توفر شبكة مستطيلة مكتملة دون نقاط مكررة.')
    source=st.radio('مصدر بيانات 3D',['مثال جاهز','إدخال يدوي','رفع CSV / Excel'],horizontal=True)
    if source=='مثال جاهز':raw=example_3d()
    elif source=='إدخال يدوي':raw=DEMO.copy()
    else:
        upload=st.file_uploader('ملف البيانات',type=['csv','xlsx'],key='3d_upload')
        if upload is None:st.stop()
        try:raw=pd.read_csv(upload) if upload.name.lower().endswith('.csv') else pd.read_excel(upload)
        except Exception as e:st.error(str(e));st.stop()
    edited=st.data_editor(raw,num_rows='dynamic',width='stretch',key='3d_editor')
    cols=list(edited.columns)
    if len(cols)<3:st.warning('يلزم 3 أعمدة.');st.stop()
    c1,c2,c3=st.columns(3)
    xc=c1.selectbox('X',cols,index=0,key='3dx')
    yc=c2.selectbox('Y',[c for c in cols if c!=xc],key='3dy')
    zc=c3.selectbox('Z',[c for c in cols if c not in [xc,yc]],key='3dz')
    df=clean(edited,[xc,yc,zc]);mode=st.selectbox('نوع الرسم',['نقاط 3D','سطح 3D','خريطة ألوان 2D'])
    if len(df)<3:st.warning('يلزم ثلاث نقاط صحيحة على الأقل.');st.stop()
    x=df[xc].to_numpy(float);y=df[yc].to_numpy(float);z=df[zc].to_numpy(float)
    xs=np.unique(x);ys=np.unique(y);complete=(len(xs)*len(ys)==len(df) and not df.duplicated([xc,yc]).any())
    if mode!='نقاط 3D' and not complete:st.warning('البيانات ليست شبكة X–Y مكتملة. اعرضيها كنقاط، أو أكملي جميع أزواج X وY دون تكرار.');st.stop()
    if mode=='خريطة ألوان 2D':
        grid=df.pivot(index=yc,columns=xc,values=zc).reindex(index=ys,columns=xs)
        fig,ax=plt.subplots(figsize=(9,6));im=ax.pcolormesh(xs,ys,grid.to_numpy(),shading='auto',cmap='viridis');fig.colorbar(im,ax=ax,label=zc);ax.set(xlabel=xc,ylabel=yc,title='PHIZKA Color Map')
    else:
        fig=plt.figure(figsize=(9,6));ax=fig.add_subplot(111,projection='3d')
        if mode=='نقاط 3D':
            artist=ax.scatter(x,y,z,c=z,cmap='viridis',s=40);fig.colorbar(artist,ax=ax,shrink=.7)
        else:
            grid=df.pivot(index=yc,columns=xc,values=zc).reindex(index=ys,columns=xs)
            X,Y=np.meshgrid(xs,ys);artist=ax.plot_surface(X,Y,grid.to_numpy(),cmap='viridis',edgecolor='none');fig.colorbar(artist,ax=ax,shrink=.7)
        ax.set(xlabel=xc,ylabel=yc,zlabel=zc,title='PHIZKA 3D')
    st.pyplot(fig);graph=png(fig);plt.close(fig)
    st.download_button('تنزيل صورة PNG',graph,'PHIZKA_3D.png','image/png')
    st.download_button('تنزيل بيانات CSV',df.to_csv(index=False).encode('utf-8-sig'),'PHIZKA_3D.csv','text/csv')
    st.caption('الرسومات ثلاثية الأبعاد هنا تُنشأ باستخدام Matplotlib، وتعمل على Streamlit Cloud دون تثبيت Gnuplot خارجي.')

else:
    st.header('📘 دليل الطالبة')
    st.markdown('''1. اختاري **قانون هوك** لتطبيق قراءات الورشة، أو **تجربة جديدة** لبياناتك الخاصة.\n2. أدخلي القراءات من الجدول أو ارفعي ملف CSV / Excel.\n3. اختاري المحور الأفقي X والرأسي Y والوحدات المناسبة.\n4. اختاري نموذج التحليل، ثم راجعي الرسم والميل وR².\n5. إذا كان لديك X وY وZ فانتقلي إلى مختبر 3D.\n6. نزّلي Word أو Excel أو CSV أو PNG.\n\n**تنبيه:** تأكدي من الوحدات والقانون النظري، ولا تفسري R² وحده بوصفه دليلًا على صحة التجربة.''')

st.divider();st.caption('PHIZKA Scientific Platform | إعداد وتطوير: البتول سالم آل سليمان')
